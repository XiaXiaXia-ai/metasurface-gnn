# -*- coding: utf-8 -*-
"""Three-slide deck: technical route (flowchart), status & directions, conference borrowables."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY = RGBColor(0x12, 0x26, 0x3A)
INK = RGBColor(0x2B, 0x38, 0x45)
GRAY = RGBColor(0x5E, 0x6B, 0x78)
MUTED = RGBColor(0x8A, 0x96, 0xA2)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
AMBER = RGBColor(0xC9, 0x7A, 0x24)
LIGHT = RGBColor(0xF4, 0xF6, 0xF8)
LINE = RGBColor(0xDD, 0xE3, 0xE9)
CREAM = RGBColor(0xFD, 0xF4, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xC9, 0xD2, 0xDA)
ARROW = RGBColor(0xAF, 0xBA, 0xC5)

FONT = "Microsoft YaHei"


def _set_font(run, size, color, bold=False, spacing=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)
    if spacing is not None:
        rPr.set("spc", str(spacing))


def _strip_style(shape):
    sp = shape._element
    style = sp.find(qn("p:style"))
    if style is not None:
        sp.remove(style)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb


def para(tf, first=False, space_before=0, space_after=0, line=1.1, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    return p


def run(p, text, size, color, bold=False, spacing=None):
    r = p.add_run()
    r.text = text
    _set_font(r, size, color, bold=bold, spacing=spacing)
    return r


def rect(slide, x, y, w, h, fill, line_color=None, shape=MSO_SHAPE.RECTANGLE, line_w=0.75):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    _strip_style(s)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line_color is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line_color
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if s.has_text_frame:
        s.text_frame.text = ""
    return s


def badge(slide, x, y, d, num, color, size=8.5):
    s = rect(slide, x, y, d, d, None, color, MSO_SHAPE.OVAL, line_w=1.0)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, num, size, color, bold=True)


def header(slide, kicker, title, subtitle, title_size=27, sub_size=10.8):
    tb = textbox(slide, 0.7, 0.38, 11.9, 0.24)
    p = para(tb.text_frame, first=True)
    run(p, "■ ", 7, TEAL)
    run(p, kicker, 9.5, TEAL, bold=True, spacing=140)

    tb = textbox(slide, 0.7, 0.64, 11.9, 0.60)
    p = para(tb.text_frame, first=True, line=1.0)
    run(p, title, title_size, NAVY, bold=True, spacing=-10)

    tb = textbox(slide, 0.7, 1.28, 11.9, 0.38)
    p = para(tb.text_frame, first=True, line=1.2)
    run(p, subtitle, sub_size, GRAY)

    rect(slide, 0.7, 1.74, 11.93, 0.012, LINE)


def pagenum(slide, idx, total=3):
    tb = textbox(slide, 11.55, 7.02, 1.08, 0.24)
    p = para(tb.text_frame, first=True, align=PP_ALIGN.RIGHT)
    run(p, f"{idx:02d}", 10, NAVY, bold=True)
    run(p, f" /{total:02d}", 8.5, MUTED)


def footnote(slide, text):
    tb = textbox(slide, 0.70, 6.96, 10.6, 0.22)
    p = para(tb.text_frame, first=True)
    run(p, text, 7.5, MUTED)


# ----------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# ============================== SLIDE 1 =====================================
s1 = prs.slides.add_slide(blank)
rect(s1, 0, 0, 13.333, 7.5, WHITE)

header(
    s1,
    "技术路线",
    "物理主干 + 学习修正 + 物理读出",
    "20×20 1-bit 编码超表面 → 8280 个方向：网络只学 400 个复权重，传播交给写死的可微算子",
    title_size=27,
)

# data bar
rect(s1, 0.85, 1.88, 11.64, 0.48, LIGHT)
rect(s1, 0.85, 1.88, 0.040, 0.48, TEAL)
tb = textbox(s1, 1.02, 1.95, 11.3, 0.34)
p = para(tb.text_frame, first=True)
run(p, "数据　", 8.6, TEAL, bold=True)
run(p, "CST 全波：20×20 · 10 GHz · 2° 网格 8280 点 · 306 s/例 · 已生成 2000 例（单频、单入射、单极化）", 8.6, INK)

# spine
cards = [
    ("01", "编码 → 阵列", ["10×10 宏比特", "block=2 → 400 单元", "每单元 1-bit 相位"]),
    ("02", "建图", ["400 节点 × 9 特征", "8 邻域 → 2964 条边", "边特征 Δx, Δy, r（÷λ）"]),
    ("03", "GNN 修正", ["3 层边条件消息传递", "Δq：400 个复数", "整图池化 → b：1 个复数"]),
    ("04", "权重合成", ["q = q0 + Δq + b", "q0 查单元库（0 参数）", "零初始化 → 起点 = M0"]),
    ("05", "物理读出", ["S = Σ qₙ·exp(jk·rₙ)", "σ = |Eθ|² + |Eφ|²", "8280 个方向 → dBsm"]),
]
CW, CG, CY, CH2 = 2.12, 0.26, 2.56, 1.56
for i, (num, name, lines) in enumerate(cards):
    x = 0.85 + i * (CW + CG)
    rect(s1, x, CY, CW, CH2, WHITE, LINE)
    rect(s1, x, CY, CW, 0.030, NAVY if i < 4 else TEAL)

    tb = textbox(s1, x + 0.13, CY + 0.12, 0.4, 0.22)
    p = para(tb.text_frame, first=True)
    run(p, num, 8, TEAL, bold=True)

    tb = textbox(s1, x + 0.13, CY + 0.33, CW - 0.22, 0.26)
    p = para(tb.text_frame, first=True)
    run(p, name, 11.5, NAVY, bold=True)

    tb = textbox(s1, x + 0.13, CY + 0.64, CW - 0.22, CH2 - 0.70)
    tf = tb.text_frame
    for j, ln in enumerate(lines):
        p = para(tf, first=(j == 0), line=1.34, space_after=3)
        run(p, ln, 8.8, GRAY)

    if i < 4:
        rect(s1, x + CW + 0.045, CY + 0.72, 0.17, 0.15, ARROW, shape=MSO_SHAPE.RIGHT_ARROW)

# data bar -> card 01
rect(s1, 1.82, 2.40, 0.15, 0.13, ARROW, shape=MSO_SHAPE.DOWN_ARROW)

# divider + caption
rect(s1, 0.85, 4.28, 11.64, 0.010, LINE)
tb = textbox(s1, 0.85, 4.34, 5.0, 0.20)
p = para(tb.text_frame, first=True)
run(p, "同一套前向，三种用法", 7.8, MUTED, bold=True)

# bottom blocks
blocks = [
    ("训练", TEAL, ["损失：dB 域 asinh 压缩 MSE", "sinθ 立体角权重 + 主瓣 ×3", "AdamW，早停只看 val"]),
    ("评估", TEAL, ["rho_lin / rho_db / bias / rmse", "分带误差 0–5° … 60–90°", "对照 M0 / M0g / M1 / M2 / M3"]),
    ("反设计（下一步）", AMBER, ["对 q 求梯度 → 松弛 + 二值化", "CST 复核 Top-K", "回写码本，形成闭环"]),
]
bw = [3.72, 3.72, 3.60]
bx = [0.85, 4.77, 8.69]
for i, (name, colr, lines) in enumerate(blocks):
    rect(s1, bx[i], 4.58, bw[i], 1.28, WHITE, LINE if colr == TEAL else AMBER)
    rect(s1, bx[i], 4.58, 0.034, 1.28, colr)

    tb = textbox(s1, bx[i] + 0.18, 4.66, bw[i] - 0.3, 0.24)
    p = para(tb.text_frame, first=True)
    run(p, name, 10, colr, bold=True)

    tb = textbox(s1, bx[i] + 0.18, 4.94, bw[i] - 0.32, 0.84)
    tf = tb.text_frame
    for j, ln in enumerate(lines):
        p = para(tf, first=(j == 0), line=1.36, space_after=2)
        run(p, ln, 8.6, INK)

# takeaway strip
rect(s1, 0.85, 6.16, 11.64, 0.46, NAVY)
tb = textbox(s1, 1.05, 6.25, 11.3, 0.30)
p = para(tb.text_frame, first=True)
run(p, "三条不变量：", 9.2, RGBColor(0x7F, 0xC4, 0xC9), bold=True)
run(p, "物理主干 0 参数　·　只学 400 个复权重　·　零初始化使 M3 起点恰为 M0", 9.2, WHITE)

footnote(s1, "依据：后续工作计划.md · GNN代码解读.md · 一份真实数据的全流程演示.md")
pagenum(s1, 1)

# ============================== SLIDE 2 =====================================
s2 = prs.slides.add_slide(blank)
rect(s2, 0, 0, 13.333, 7.5, WHITE)

header(
    s2,
    "现状与方向",
    "镜面已补上，误差迁移到大角域",
    "5 epoch / 291 例快照 / 44 例 test 的真实数字：骨架已经闭合，剩下的是四个维度的泛化",
    title_size=27,
)

# three stat tiles
stats = [
    ("镜面误差  θ = 0", "−16.1 dB  →  +0.16 dB", "M0  →  M3"),
    ("RMSE", "4.65 dB  →  2.64 dB", "全角域"),
    ("rho_db", "0.865  →  0.958", "弥散结构"),
]
for i, (lab, val, sub) in enumerate(stats):
    x = 0.70 + i * 2.03
    rect(s2, x, 1.92, 1.88, 0.94, LIGHT)
    tb = textbox(s2, x + 0.14, 1.99, 1.62, 0.22)
    p = para(tb.text_frame, first=True)
    run(p, lab, 8.2, GRAY)
    tb = textbox(s2, x + 0.14, 2.22, 1.66, 0.26)
    p = para(tb.text_frame, first=True)
    run(p, val, 10.2, NAVY, bold=True)
    tb = textbox(s2, x + 0.14, 2.52, 1.62, 0.22)
    p = para(tb.text_frame, first=True)
    run(p, sub, 7.6, MUTED)

# banded error bars
tb = textbox(s2, 0.70, 3.10, 5.9, 0.24)
p = para(tb.text_frame, first=True)
run(p, "M3 分带 RMSE（test）", 9, NAVY, bold=True)
run(p, "　　误差已从镜面锥迁到 60–90° 掠射区", 8, MUTED)

bands = [("0–5°", 0.29), ("5–15°", 0.98), ("15–30°", 0.99), ("30–60°", 1.50), ("60–90°", 4.10)]
maxv, maxw = 4.10, 3.05
for i, (lab, v) in enumerate(bands):
    y = 3.42 + i * 0.38
    tb = textbox(s2, 0.70, y - 0.02, 0.75, 0.22)
    p = para(tb.text_frame, first=True)
    run(p, lab, 8.2, GRAY)
    w = max(0.06, v / maxv * maxw)
    rect(s2, 1.50, y + 0.02, w, 0.15, AMBER if v > 2 else TEAL)
    tb = textbox(s2, 1.50 + w + 0.08, y - 0.02, 0.7, 0.22)
    p = para(tb.text_frame, first=True)
    run(p, f"{v:.2f} dB", 8.2, INK)

tb = textbox(s2, 0.70, 5.42, 6.0, 0.45)
p = para(tb.text_frame, first=True, line=1.25)
run(p, "M0 的 0–5° 误差是 14.0 dB，M3 压到 0.29 dB；", 8.2, INK)
run(p, "新的短板是大角域（4.10 dB）——倾斜因子与边缘效应。", 8.2, INK)

# directions
tb = textbox(s2, 6.95, 1.92, 5.7, 0.24)
p = para(tb.text_frame, first=True)
run(p, "往哪走：四个维度", 9, NAVY, bold=True)

dirs = [
    ("01", "目标 · 换挡", "随机码 → 结构化码本（斜坡 / 多波束 / 零陷）", "预测 RCS → 匹配目标方向图"),
    ("02", "规模 · 外推", "固定 20×20 → 可变口径（滑窗 / 分块）", "GNN 与尺寸无关，算子按尺寸预计算"),
    ("03", "频率 · 宽带", "单频 10 GHz → 多频", "阵中 f0 = 9.6 GHz、Q = 4.86，色散进单元库"),
    ("04", "单元 · 异构", "双半径贴片 → 异形单元混排", "图相对 CNN 的真正价值；M4 加长程边"),
]
for i, (num, head, now, nxt) in enumerate(dirs):
    y = 2.20 + i * 0.98
    rect(s2, 6.95, y, 5.68, 0.90, WHITE, LINE)
    rect(s2, 6.95, y, 0.034, 0.90, TEAL if i != 0 else AMBER)

    badge(s2, 7.10, y + 0.10, 0.28, num, TEAL, size=7.6)

    tb = textbox(s2, 7.48, y + 0.09, 2.0, 0.24)
    p = para(tb.text_frame, first=True)
    run(p, head, 10.5, NAVY, bold=True)

    tb = textbox(s2, 7.48, y + 0.37, 5.0, 0.22)
    p = para(tb.text_frame, first=True)
    run(p, now, 7.9, GRAY)

    tb = textbox(s2, 7.48, y + 0.61, 5.0, 0.22)
    p = para(tb.text_frame, first=True)
    run(p, nxt, 7.9, INK)

# bottom band
rect(s2, 0.70, 6.14, 11.93, 0.44, NAVY)
tb = textbox(s2, 0.92, 6.22, 11.5, 0.28)
p = para(tb.text_frame, first=True)
run(p, "判断：", 9.2, RGBColor(0x7F, 0xC4, 0xC9), bold=True)
run(p, "方法骨架已经闭合，接下来不是换网络，而是把「目标 × 规模 × 频率 × 单元」四个维度逐个打开。", 9.2, WHITE)

footnote(s2, "数字来源：demo_out/eval/metrics_summary.csv（M3 与 M0，test，44 例）")
pagenum(s2, 2)

# ============================== SLIDE 3 =====================================
s3 = prs.slides.add_slide(blank)
rect(s3, 0, 0, 13.333, 7.5, WHITE)

header(
    s3,
    "会议借鉴",
    "两篇会议工作，四件可以借的事",
    "只取组织方式——主干、残差、读出、闭环；网络结构不必照搬",
    title_size=27,
)

# column heads
for (x, w, txt) in [(1.32, 2.50, "借鉴点"), (4.02, 3.40, "他们怎么做的"), (7.72, 4.80, "我们怎么用")]:
    tb = textbox(s3, x, 1.86, w, 0.22)
    p = para(tb.text_frame, first=True)
    run(p, txt, 8.2, MUTED, bold=True, spacing=60)

rows = [
    ("01", "边按机制分，不按距离分",
     "LoS / 反射 / 绕射三类边；路径限最多 2 次反射、1 次绕射",
     "M4 的长程边照此办理：近场 / 表面波 / 衍射各设截断，而不是把 k-NN 半径调大"),
    ("02", "主干 × 可退化残差",
     "m = A·fθ(h)，令 fθ ≡ 1 即退回解析解，用作消融基线",
     "我们已是 q0 + Δq + b：δ ≡ 0 就是 M0。写进方法论，并当作消融与 OOD 兜底"),
    ("03", "物理读出，不学 decoder",
     "按物理路径池化复数消息，再取 dB；不是全图平均",
     "8280×400 传播算子直接读出；消息聚合用 sum，而不是定长邻居拼接"),
    ("04", "闭环 rollout",
     "0 → 100 m 逐面推进，训练时喂自己的预测压住漂移",
     "尺寸 / 频率外推时用 scheduled sampling 抑制误差累积（提议，尚未实现）"),
]
ROW_H, STEP, Y0 = 0.90, 0.98, 2.16
for i, (num, head, theirs, ours) in enumerate(rows):
    y = Y0 + i * STEP
    if i % 2 == 0:
        rect(s3, 0.70, y, 11.93, ROW_H, LIGHT)

    tb = textbox(s3, 0.84, y, 0.45, ROW_H, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tb.text_frame, first=True)
    run(p, num, 11.5, PALE, bold=True)

    tb = textbox(s3, 1.32, y, 2.50, ROW_H, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tb.text_frame, first=True, line=1.12)
    run(p, head, 11.2, NAVY, bold=True)

    tb = textbox(s3, 4.02, y, 3.40, ROW_H, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tb.text_frame, first=True, line=1.28)
    run(p, theirs, 8.8, GRAY)

    tb = textbox(s3, 7.72, y, 4.80, ROW_H, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tb.text_frame, first=True, line=1.28)
    run(p, ours, 8.8, INK)

rect(s3, 0.70, 6.30, 11.93, 0.012, LINE)
tb = textbox(s3, 0.70, 6.42, 11.9, 0.28)
p = para(tb.text_frame, first=True)
run(p, "一句话：", 8.8, AMBER, bold=True)
run(p, "借的是「物理负责外推、网络只补残差」的组织方式，不是他们的网络结构。", 8.8, INK)

pagenum(s3, 3)

out = r"D:\DSHNESS\GNN\GNN_编码超表面_技术路线_三页.pptx"
prs.save(out)
print("saved", out)
