# -*- coding: utf-8 -*-
"""Build the 2-slide deck: lessons from two ISAP talks -> our coding-metasurface GNN (M3)."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY = RGBColor(0x12, 0x26, 0x3A)
INK = RGBColor(0x2B, 0x38, 0x45)
GRAY = RGBColor(0x5E, 0x6B, 0x78)
MUTED = RGBColor(0x8A, 0x96, 0xA2)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
AMBER = RGBColor(0xC9, 0x7A, 0x24)
LIGHT = RGBColor(0xF4, 0xF6, 0xF8)
LINE = RGBColor(0xDD, 0xE3, 0xE9)
CREAM = RGBColor(0xFD, 0xF4, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xC9, 0xD2, 0xDA)
ARROW = RGBColor(0xAF, 0xBA, 0xC5)

FONT = "Microsoft YaHei"


def _set_font(run, size, color, bold=False, italic=False, spacing=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)
    if spacing is not None:
        rPr.set("spc", str(spacing))


def _strip_style(shape):
    sp = shape._element
    style = sp.find(qn("p:style"))
    if style is not None:
        sp.remove(style)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb


def para(tf, first=False, space_before=0, space_after=0, line=1.1, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    return p


def run(p, text, size, color, bold=False, spacing=None):
    r = p.add_run()
    r.text = text
    _set_font(r, size, color, bold=bold, spacing=spacing)
    return r


def rect(slide, x, y, w, h, fill, line_color=None, shape=MSO_SHAPE.RECTANGLE, line_w=0.75):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    _strip_style(s)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line_color is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line_color
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if s.has_text_frame:
        s.text_frame.text = ""
    return s


def badge(slide, x, y, d, num, color):
    s = rect(slide, x, y, d, d, None, color, MSO_SHAPE.OVAL, line_w=1.0)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, num, 8.5, color, bold=True)


def header(slide, kicker, title, subtitle, title_size=26, sub_size=10.8):
    tb = textbox(slide, 0.7, 0.38, 11.9, 0.24)
    p = para(tb.text_frame, first=True)
    run(p, "■ ", 7, TEAL)
    run(p, kicker, 9.5, TEAL, bold=True, spacing=140)

    tb = textbox(slide, 0.7, 0.64, 11.9, 0.60)
    p = para(tb.text_frame, first=True, line=1.0)
    run(p, title, title_size, NAVY, bold=True, spacing=-10)

    tb = textbox(slide, 0.7, 1.28, 11.9, 0.38)
    p = para(tb.text_frame, first=True, line=1.2)
    run(p, subtitle, sub_size, GRAY)

    rect(slide, 0.7, 1.74, 11.93, 0.012, LINE)


def pagenum(slide, idx):
    tb = textbox(slide, 11.55, 7.02, 1.08, 0.24)
    p = para(tb.text_frame, first=True, align=PP_ALIGN.RIGHT)
    run(p, f"{idx:02d}", 10, NAVY, bold=True)
    run(p, " /02", 8.5, MUTED)


# ----------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# ============================== SLIDE 1 =====================================
s1 = prs.slides.add_slide(blank)
rect(s1, 0, 0, 13.333, 7.5, WHITE)

header(
    s1,
    "ISAP · 同一课题组的两条技术路线",
    "两篇会议工作，四个可直接落到 M3 的设计决策",
    "传播图 GNN（室内路径损耗）× 物理信息自回归网络（隧道 RSS）：两篇都是「解析主干 + 学习修正」，与我们的 γ = Γ_bit(1+δ) 同构",
)

cards = [
    (
        "01", "边按机制分，不按距离分",
        "LoS·反射·绕射三类边；路径限最多两次反射、一次绕射",
        "M4 要加的「长程 / 多尺度边」同样需要「哪种机制、连多远」的规则",
        "近场 / 表面波 / 衍射分类建边、各自截断；边缘单元加边界编码",
    ),
    (
        "02", "主干 × 可退化残差",
        "m = A·fθ(h)；令 fθ ≡ 1 即退回解析解，用作消融基线",
        "我们的 γ = Γ_bit(1+δ)，δ ≡ 0 正是 M0",
        "把 δ≡0 写成结构约束 + 消融 + OOD 兜底；δ 场直接当耦合图看",
    ),
    (
        "03", "物理读出，不学 decoder",
        "E = Σ m 再取 dB；池化按物理路径，不按全图节点",
        "远场本就是 Σγₙ·exp(jk·rₙ) 的线性泛函",
        "读出层直接用 8280×400 传播算子；聚合用 sum/attention 而非定长拼接",
    ),
    (
        "04", "闭环 rollout",
        "0→100 m 逐面推进，训练时喂自己的预测压住漂移",
        "20×20 → 更大口径的分块外推、f0 → 相邻频点，都是 rollout",
        "块间 / 频间自回归 + scheduled sampling，攻「可变规模 / 多频段」",
    ),
]

CW, CH, GX, GY = 5.82, 2.04, 0.29, 0.24
for i, (num, head, theirs, mapping, ours) in enumerate(cards):
    col, row = i % 2, i // 2
    x = 0.70 + col * (CW + GX)
    y = 1.94 + row * (CH + GY)

    rect(s1, x, y, CW, CH, LIGHT)
    rect(s1, x, y, 0.042, CH, TEAL)
    badge(s1, x + 0.22, y + 0.19, 0.32, num, TEAL)

    tb = textbox(s1, x + 0.66, y + 0.17, CW - 0.9, 0.34)
    p = para(tb.text_frame, first=True)
    run(p, head, 13.5, NAVY, bold=True, spacing=-5)

    tb = textbox(s1, x + 0.24, y + 0.52, CW - 0.48, CH - 0.60, anchor=MSO_ANCHOR.MIDDLE)
    tf = tb.text_frame
    rows = [("他们　", theirs, GRAY), ("同构　", mapping, AMBER), ("我们　", ours, TEAL)]
    for j, (lab, txt, colr) in enumerate(rows):
        p = para(tf, first=(j == 0), line=1.42, space_after=8)
        run(p, lab, 9.5, colr, bold=True)
        run(p, txt, 10.5, INK)

# takeaway band
band_y = 6.38
rect(s1, 0.70, band_y, 11.93, 0.44, NAVY)
tb = textbox(s1, 0.94, band_y + 0.10, 11.5, 0.26)
p = para(tb.text_frame, first=True)
run(p, "共同点：", 9.5, RGBColor(0x7F, 0xC4, 0xC9), bold=True)
run(p, "物理算子负责外推，网络只补 Γ_bit 算不准的那部分——这正是 M3 相对 M0/M1/M2 的立身之本。", 9.5, WHITE)

tb = textbox(s1, 0.70, 6.96, 11.2, 0.22)
p = para(tb.text_frame, first=True)
run(p, "原文对应：propagation graph　·　analytic route messages　·　graph residual & field synthesis　·　autoregressive prediction & training", 7.5, MUTED)

pagenum(s1, 1)

# ============================== SLIDE 2 =====================================
s2 = prs.slides.add_slide(blank)
rect(s2, 0, 0, 13.333, 7.5, WHITE)

header(
    s2,
    "落地方案",
    "落到 M3：五步架构，三条改造，一条评估红线",
    "把会议里的「物理主干 + 学习残差」直接写进 20×20 编码超表面的 γ / δ 路线（10 GHz，逐单元 1-bit）",
    title_size=26,
)

steps = [
    ("1", "图构建", "节点：几何·码位·边界", "边：近场/表面波/衍射，分类截断"),
    ("2", "物理主干", "γ = Γ_bit（δ ≡ 0）", "传播算子 A（8280×400）预计算"),
    ("3", "GNN 修正", "输出 δ（复数，400 DOF）", "sum / attention 聚合"),
    ("4", "物理读出", "E = Σ γₙ·exp(jk·rₙ)", "不学 decoder"),
    ("5", "双头输出", "方向图（dB）", "＋ 绝对增益（不归一化）"),
]

BW, BG, BH = 2.16, 0.245, 1.02
for i, (num, name, l1, l2) in enumerate(steps):
    x = 0.72 + i * (BW + BG)
    rect(s2, x, 1.96, BW, BH, WHITE, LINE)
    rect(s2, x, 1.96, BW, 0.032, NAVY if i < 4 else TEAL)

    badge(s2, x + 0.11, 2.08, 0.27, num, TEAL)

    tb = textbox(s2, x + 0.46, 2.07, BW - 0.55, 0.28)
    p = para(tb.text_frame, first=True)
    run(p, name, 11.5, NAVY, bold=True)

    tb = textbox(s2, x + 0.12, 2.46, BW - 0.20, 0.22)
    p = para(tb.text_frame, first=True)
    run(p, l1, 7.8, GRAY)

    tb = textbox(s2, x + 0.12, 2.70, BW - 0.20, 0.22)
    p = para(tb.text_frame, first=True)
    run(p, l2, 7.8, GRAY)

    if i < 4:
        rect(s2, x + BW + 0.052, 2.38, 0.14, 0.18, ARROW, shape=MSO_SHAPE.RIGHT_ARROW)

rationales = [
    ("01", "边按机制分，不按距离分", TEAL,
     "「最多 2 次反射、1 次绕射」可直接照搬：每类耦合独立截断半径与阶数。M4 的长程边因此有物理依据，而不是把 k-NN 半径调大。"),
    ("02", "δ ≡ 0 必须能退化", TEAL,
     "fθ≡1 是人家的消融基线，我们的 M0 就是 δ≡0。训练给 |δ| 加正则、推理留 δ=0 兜底，OOD 至少不劣于物理下限。"),
    ("03", "尺寸 / 频率维 rollout", AMBER,
     "他们把 100 步 rollout 训稳；我们用在更大口径的分块外推与 f0→相邻频点，让误差不随块数、频差累积。"),
]
CW3, GAP3 = 3.72, 0.38
for i, (num, head, colr, body) in enumerate(rationales):
    x = 0.72 + i * (CW3 + GAP3)
    rect(s2, x, 3.44, 0.034, 1.62, colr)

    tb = textbox(s2, x + 0.18, 3.40, 0.7, 0.32)
    p = para(tb.text_frame, first=True)
    run(p, num, 16, PALE, bold=True)

    tb = textbox(s2, x + 0.18, 3.80, CW3 - 0.2, 0.28)
    p = para(tb.text_frame, first=True)
    run(p, head, 12.5, NAVY, bold=True)

    tb = textbox(s2, x + 0.18, 4.14, CW3 - 0.22, 0.90)
    p = para(tb.text_frame, first=True, line=1.24)
    run(p, body, 9.5, INK)

# red-line band
rect(s2, 0.72, 5.30, 11.90, 0.84, CREAM)
rect(s2, 0.72, 5.30, 0.042, 0.84, AMBER)

tb = textbox(s2, 0.94, 5.37, 1.2, 0.22)
p = para(tb.text_frame, first=True)
run(p, "评估红线", 8.5, AMBER, bold=True, spacing=60)

tb = textbox(s2, 0.94, 5.62, 11.45, 0.44)
p = para(tb.text_frame, first=True, line=1.15)
run(p, "报方向图指标（镜面 RMSE / 旁瓣 / 零深），不报归一化场 MSE", 9.6, NAVY, bold=True)
run(p, "　·　数据生成 / 图构建 / 推理 分项计时", 9.6, INK)
run(p, "　·　跨规模、跨频、跨单元库 OOD，多点多例", 9.6, INK)

tb = textbox(s2, 0.72, 6.26, 11.90, 0.34)
p = para(tb.text_frame, first=True, line=1.15)
run(p, "反面教材：", 8.3, AMBER, bold=True)
run(p, "他们自述的局限——证据只有一点一几何；≈2λ 采样分辨不了驻波；逐面归一化丢掉绝对值。对应我们的三个坑：单 case 报喜、角度采样不足、归一化 MSE 跨任务比较。", 8.3, MUTED)

rect(s2, 0.70, 6.80, 11.93, 0.012, LINE)
tb = textbox(s2, 0.70, 6.90, 11.0, 0.24)
p = para(tb.text_frame, first=True)
run(p, "依据：想法笔记.md、后续工作计划.md（M0–M4、γ = Γ_bit(1+δ)、8280×400 传播算子）", 7.8, MUTED)

pagenum(s2, 2)

out = r"D:\DSHNESS\GNN\GNN_超表面阵列_会议借鉴_两页.pptx"
prs.save(out)
print("saved", out)
